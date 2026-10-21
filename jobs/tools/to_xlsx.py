#!/usr/bin/env python3
"""CSV (UTF-8) -> μορφοποιημένο .xlsx για τον Δήμο.
Χρήση: python3 to_xlsx.py input.csv output.xlsx "Τίτλος φύλλου" [--approval "Έγκριση"]
- Bold header, freeze, wrap, πλάτη στηλών, links κλικαριστά.
- Στήλη έγκρισης: dropdown ✅ Ναι / ✏️ Αλλαγή / ❌ Όχι.
"""
import csv, sys, argparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

p = argparse.ArgumentParser()
p.add_argument("src"); p.add_argument("dst"); p.add_argument("title")
p.add_argument("--approval", default=None, help="όνομα στήλης έγκρισης")
a = p.parse_args()

rows = list(csv.reader(open(a.src, encoding="utf-8")))
head, body = rows[0], rows[1:]
wb = Workbook(); ws = wb.active; ws.title = a.title[:31]
ws.append(head)
for r in body: ws.append(r)

hdr_fill = PatternFill("solid", fgColor="1F3A2E")
thin = Side(style="thin", color="D9D9D9")
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = hdr_fill
    c.alignment = Alignment(vertical="center", wrap_text=True)
ws.row_dimensions[1].height = 30
ws.freeze_panes = "B2"

for i, h in enumerate(head, start=1):
    col = get_column_letter(i)
    longest = max([len(h)] + [len(str(r[i-1])) for r in body if len(r) >= i] or [10])
    ws.column_dimensions[col].width = max(10, min(55, longest * 0.9 + 2))
    for cell in ws[col][1:]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
        v = str(cell.value or "")
        if v.startswith("http://") or v.startswith("https://"):
            cell.hyperlink = v; cell.font = Font(color="1155CC", underline="single")

if a.approval and a.approval in head:
    idx = head.index(a.approval) + 1; col = get_column_letter(idx)
    dv = DataValidation(type="list", formula1='"✅ Ναι,✏️ Αλλαγή,❌ Όχι"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"{col}2:{col}{len(body)+1}")
    ws.column_dimensions[col].width = 14
    for cell in ws[col][1:]:
        cell.fill = PatternFill("solid", fgColor="FFF8E1")

ws.auto_filter.ref = ws.dimensions
wb.save(a.dst)
print(f"ok {a.dst}: {len(body)} γραμμές")
